#!/usr/bin/env python3
"""
Adds the phase-1 buying sheet to sub-mission_parts.xlsx.

PARTS_LIST.csv carries a `phase` column:
  1    = buy now       2 = later phase       cut = dead line, do not buy

The sheet is grouped by shop (that is how it gets ordered). The CSV stays the
machine source; this workbook is the document you read and hand over.
Shipping is not in the CSV, so it is estimated at the bottom.

Run after build_parts_xlsx.py:
  PYTHONPATH=~/.local/lib/python3.14/site-packages /usr/bin/python3 tools/build_order_v1.py
"""
import csv
import math
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

XLSX = "/home/filip/sub-mission/docs/sub-mission_parts.xlsx"
CSV = "/home/filip/sub-mission/docs/PARTS_LIST.csv"
SHEET = "Order v1 (now)"
SHIPPING_EST = 25.0

SHOPS = [("tme.eu", "TME"), ("sensorsone", "sensorsone (quote)"),
         ("techfun", "techfun.sk"), ("prusa", "Prusa3D"),
         ("x-teamrc", "x-teamrc.com"), ("tmotor", "T-Motor"),
         ("magsy", "magsy.sk"), ("supermagnet", "supermagneticshop"),
         ("neven", "neven.cz"), ("nerezmaterial", "nerezmaterial.cz"),
         ("laskakit", "laskakit.cz"), ("jlcpcb", "JLCPCB"),
         ("local", "local / shop"), ("printed", "printed in-house"),
         ("in CAD", "in CAD")]

DARK, MID, LIGHT, TOTAL = "1F3864", "2E5C8A", "EAF0F8", "C6E0B4"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0.00'


def n(x):
    try:
        return float(x)
    except Exception:
        return 0.0


def shop_of(src):
    s = (src or "").lower()
    for key, name in SHOPS:
        if key in s:
            return name
    return "local / shop"


def qty_of(r):
    return sum(n(r.get(k, 0)) for k in ("qty_probe", "qty_site", "qty_gateway"))


def cost_of(r):
    """Owned items stay on the list but cost nothing."""
    if (r.get("owned") or "").upper() == "YES":
        return 0.0
    return qty_of(r) * n(r["price_eur"])


def height_for(text, width_chars, base=13.5):
    lines = max(1, math.ceil(len(str(text or "")) / max(width_chars, 8)))
    return base * lines + 2


def main():
    rows = list(csv.DictReader(open(CSV, encoding="utf-8")))
    p1 = [r for r in rows if r.get("phase") == "1"]
    by_shop = defaultdict(list)
    for r in p1:
        by_shop[shop_of(r["source"])].append(r)

    goods = sum(cost_of(r) for r in p1)

    wb = load_workbook(XLSX)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "sub-mission - phase 1 order"
    c.font = Font(bold=True, size=16, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=DARK)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = ("Grouped by shop, because that is how the order is placed. No spares added. "
               "Shipping is not in the parts list and is estimated at the bottom.")
    c.font = Font(italic=True, size=9, color="404040")
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    ws.merge_cells("A3:F3")
    c = ws["A3"]
    c.value = ("%d lines      goods %.2f EUR      shipping est. %.2f EUR      "
               "LANDED %.2f EUR" % (len(p1), goods, SHIPPING_EST, goods + SHIPPING_EST))
    c.font = Font(bold=True, size=11, color=DARK)
    c.fill = PatternFill("solid", fgColor=TOTAL)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[3].height = 26

    row = 5
    for sh in sorted(by_shop, key=lambda k: -sum(cost_of(x) for x in by_shop[k])):
        items = sorted(by_shop[sh], key=lambda x: -cost_of(x))
        sub = sum(cost_of(x) for x in items)

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=sh.upper())
        cell.font = Font(bold=True, size=11, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=MID)
        cell.alignment = Alignment(vertical="center", indent=1)
        tot = ws.cell(row=row, column=6, value=round(sub, 2))
        tot.number_format = MONEY
        tot.font = Font(bold=True, size=11, color="FFFFFF")
        tot.fill = PatternFill("solid", fgColor=MID)
        tot.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1

        for col, h in enumerate(["ID", "Part", "Specification / note", "Qty",
                                 "Unit EUR", "Total EUR"], start=1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = Font(bold=True, size=9, color=DARK)
            cell.fill = PatternFill("solid", fgColor=LIGHT)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center",
                                       horizontal="right" if col > 3 else "left")
        row += 1

        for i, x in enumerate(items):
            q, p = qty_of(x), n(x["price_eur"])
            owned = (x.get("owned") or "").upper() == "YES"
            spec = x["spec"]
            if x["notes"]:
                spec = "%s  -  %s" % (spec, x["notes"])
            if owned:
                spec = "MAM  |  " + spec
            vals = [x["id"], x["part"], spec, int(q), p, round(cost_of(x), 2)]
            ws.row_dimensions[row].height = height_for(spec, 74)
            for col, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.border = BORDER
                cell.font = Font(size=9)
                cell.alignment = Alignment(vertical="top", wrap_text=(col == 3),
                                           horizontal="right" if col > 3 else "left")
                if col in (5, 6):
                    cell.number_format = MONEY
                if owned:
                    cell.fill = PatternFill("solid", fgColor="E2EFDA")
                elif i % 2:
                    cell.fill = PatternFill("solid", fgColor="F7F9FC")
            row += 1
        row += 1

    for label, val, strong in (("GOODS", goods, False),
                               ("shipping, estimated", SHIPPING_EST, False),
                               ("LANDED TOTAL", goods + SHIPPING_EST, True)):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=label)
        cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        cell.font = Font(bold=True, size=12 if strong else 10)
        cell.fill = PatternFill("solid", fgColor=TOTAL if strong else "F2F2F2")
        cell.border = BORDER
        tot = ws.cell(row=row, column=6, value=round(val, 2))
        tot.number_format = MONEY
        tot.font = Font(bold=True, size=12 if strong else 10)
        tot.fill = PatternFill("solid", fgColor=TOTAL if strong else "F2F2F2")
        tot.alignment = Alignment(horizontal="right", vertical="center")
        tot.border = BORDER
        ws.row_dimensions[row].height = 22 if strong else 18
        row += 1

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    cell = ws.cell(row=row, column=1, value="NOT IN THIS ORDER")
    cell.font = Font(bold=True, size=12, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=DARK)
    cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 22
    row += 1

    for phase, title, colr in (("2", "phase 2 - buy later", "7F7F7F"),
                               ("cut", "cut - dead line, do not buy", "C00000")):
        cell = ws.cell(row=row, column=1, value=title.upper())
        cell.font = Font(bold=True, size=10, color=colr)
        row += 1
        for r in [x for x in rows if x.get("phase") == phase]:
            price = r["price_eur"] if r["price_eur"] not in ("", "-") else "-"
            ws.cell(row=row, column=1, value=r["id"]).font = Font(size=9)
            ws.cell(row=row, column=2, value=r["part"][:52]).font = Font(size=9)
            ws.cell(row=row, column=3, value=r["notes"][:150]).font = Font(size=8, italic=True)
            cell = ws.cell(row=row, column=5, value=price)
            cell.font = Font(size=9, color=colr)
            cell.alignment = Alignment(horizontal="right")
            ws.row_dimensions[row].height = max(13, height_for(r["notes"], 105))
            row += 1
        row += 1

    for col, w in zip("ABCDEF", [7, 26, 76, 5, 10, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5

    wb.save(XLSX)
    print("phase 1: %d lines, goods %.2f + shipping %.2f = %.2f EUR"
          % (len(p1), goods, SHIPPING_EST, goods + SHIPPING_EST))


if __name__ == "__main__":
    main()
